from datetime import datetime
from json import dumps
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source" / "AZ-GOLF-2026-Overseed-Schedule.xlsx"
OUTPUT = ROOT / "data" / "courses.js"

# Cities and communities in Maricopa County represented in the source workbook.
# Queen Creek ZIP 85142 is in the Maricopa County portion; 85140 is excluded.
MARICOPA_CITIES = {
    "Anthem", "Avondale", "Buckeye", "Carefree", "Cave Creek", "Chandler",
    "El Mirage", "Fort McDowell", "Fountain Hills", "Gilbert", "Glendale",
    "Goodyear", "Laveen Village", "Litchfield Park", "Mesa", "Paradise Valley",
    "Peoria", "Phoenix", "Queen Creek", "Rio Verde", "Scottsdale", "Sun City",
    "Sun City West", "Sun Lakes", "Surprise", "Tempe", "Waddell", "Wickenburg",
}


def clean(value):
    return str(value).strip() if value is not None else ""


def iso(value):
    return value.strftime("%Y-%m-%d") if isinstance(value, datetime) else None


def included(city, zipcode):
    if city not in MARICOPA_CITIES:
        return False
    if city == "Queen Creek" and str(zipcode) != "85142":
        return False
    return True


def main():
    ws = load_workbook(SOURCE, data_only=True, read_only=True).active
    courses = []
    for facility, course, overseed, closure, reopening, city, zipcode in ws.iter_rows(
        min_row=2, values_only=True
    ):
        facility = clean(facility)
        city = clean(city)
        if not facility or not included(city, zipcode):
            continue
        courses.append({
            "facility": facility,
            "course": clean(course) or None,
            "overseed": clean(overseed).lower() == "yes",
            "closure": iso(closure),
            "reopening": iso(reopening),
            "city": city,
            "zip": str(zipcode).strip() if zipcode is not None else "",
        })

    courses.sort(key=lambda item: (item["facility"].casefold(), (item["course"] or "").casefold()))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    payload = dumps(courses, ensure_ascii=False, separators=(",", ":"))
    OUTPUT.write_text(
        "// Generated from the Arizona Golf Association 2026 overseed schedule.\n"
        f"window.OCTOBER_GOLF_COURSES={payload};\n",
        encoding="utf-8",
    )
    print(f"Wrote {len(courses)} Maricopa County course records to {OUTPUT}")


if __name__ == "__main__":
    main()
